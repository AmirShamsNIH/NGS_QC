#!/usr/bin/env python3
"""
collect_qc_logs.py – NGS-QC post-run log collector
====================================================

Reads ``qc_manifest.json`` (written by run_ngs_qc.py) for the list of every
scheduled sample × tool job, checks each job's log and expected outputs,
and produces a multi-sheet XLSX execution report.

Sheet layout
------------
  Summary       Colour-coded sample × tool matrix (PASS / FAIL / MISSING / N/A)
  Failed Jobs   Only the failed/missing entries with log snippets
  <one per tool> All samples – status + outputs + log tail

Usage
-----
::

    python collect_qc_logs.py STUDY_DIR STUDY_NAME [--output REPORT.xlsx]

    STUDY_DIR   The per-study output directory created by run_ngs_qc.py,
                e.g. /data/output/MyStudy/
    STUDY_NAME  Study label (used in the report title)
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit(
        "openpyxl is required: pip install openpyxl"
    )

# ---------------------------------------------------------------------------
# Failure detection
# ---------------------------------------------------------------------------

MANIFEST_NAME = "qc_manifest.json"

# Patterns in log content that indicate a catastrophic failure
_FATAL_PATTERNS: list[re.Pattern] = [
    re.compile(r"Traceback \(most recent call last\)"),
    re.compile(r"\bKilled\b"),
    re.compile(r"Segmentation fault", re.IGNORECASE),
    re.compile(r"Out of memory", re.IGNORECASE),
    re.compile(r"slurmstepd: error:", re.IGNORECASE),
    re.compile(r"command not found", re.IGNORECASE),
    re.compile(r"No such file or directory"),
    re.compile(r"\bRuntimeError\b"),
    re.compile(r"java\.lang\.\w*(Error|Exception)"),
]

_LOG_TAIL_LINES = 30   # lines to include in per-tool detail sheets


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

@dataclass
class ToolResult:
    status: str                  # "PASS", "FAIL", "MISSING" or "N/A"
    log_path: str
    log_exists: bool
    log_tail: str                # last N lines of the log
    output_files: list[str] = field(default_factory=list)
    missing_outputs: list[str] = field(default_factory=list)
    fatal_lines: list[str]  = field(default_factory=list)


_NOT_APPLICABLE = ToolResult("N/A", "not scheduled for this sample", False, "")


# ---------------------------------------------------------------------------
# Core logic
# ---------------------------------------------------------------------------

def load_manifest(study_dir: Path) -> dict[str, Any]:
    path = study_dir / MANIFEST_NAME
    if not path.is_file():
        sys.exit(
            f"ERROR: {path} not found. It is written by run_ngs_qc.py; "
            "re-generate the scripts for this study."
        )
    return json.loads(path.read_text())


def _fatal_lines(text: str) -> list[str]:
    return [
        line.strip()
        for line in text.splitlines()
        if any(p.search(line) for p in _FATAL_PATTERNS)
    ]


def _existing_outputs(pattern: str) -> list[str]:
    """Non-empty files matching *pattern* (a path, or a glob for AfterQC)."""
    paths = glob.glob(pattern) if glob.has_magic(pattern) else [pattern]
    return [p for p in paths if os.path.isfile(p) and os.path.getsize(p) > 0]


def check_job(job: dict[str, Any]) -> ToolResult:
    """Determine pass / fail / missing for one scheduled sample × tool job.

    PASS requires every expected output to exist and be non-empty and no
    fatal pattern in the log.
    """
    log_path   = Path(job["log"])
    log_exists = log_path.is_file()
    # strip control chars (progress bars / ANSI codes) that Excel rejects
    log_text   = ILLEGAL_CHARACTERS_RE.sub(
        "", log_path.read_text(errors="replace")) if log_exists else ""
    log_tail   = "\n".join(log_text.splitlines()[-_LOG_TAIL_LINES:])

    found, missing = [], []
    for pattern in job["outputs"]:
        hits = _existing_outputs(pattern)
        found += hits
        if not hits:
            missing.append(pattern)

    fatal = _fatal_lines(log_text)

    if not log_exists and not found:
        status = "MISSING"
    elif not missing and not fatal:
        status = "PASS"
    else:
        status = "FAIL"

    return ToolResult(
        status          = status,
        log_path        = str(log_path) if log_exists else "NOT FOUND",
        log_exists      = log_exists,
        log_tail        = log_tail,
        output_files    = found,
        missing_outputs = missing,
        fatal_lines     = fatal,
    )


def collect(manifest: dict[str, Any]) -> dict[str, dict[str, ToolResult]]:
    """Return nested dict: results[sample][tool_name] = ToolResult.

    Every sample discovered at generation time is listed, so a sample whose
    jobs never ran shows up as MISSING rather than disappearing.
    """
    results: dict[str, dict[str, ToolResult]] = {
        s["name"]: {t: _NOT_APPLICABLE for t in manifest["tools"]}
        for s in manifest["samples"]
    }
    for job in manifest["jobs"]:
        results[job["sample"]][job["tool"]] = check_job(job)
    return results


# ---------------------------------------------------------------------------
# XLSX colours / styles
# ---------------------------------------------------------------------------

_FILL_PASS    = PatternFill("solid", fgColor="92D050")   # green
_FILL_FAIL    = PatternFill("solid", fgColor="FF4B4B")   # red
_FILL_MISSING = PatternFill("solid", fgColor="FFC000")   # amber
_FILL_NA      = PatternFill("solid", fgColor="D9D9D9")   # grey
_FILL_HEADER  = PatternFill("solid", fgColor="4472C4")   # blue
_FILL_TITLE   = PatternFill("solid", fgColor="1F3864")   # dark navy
_FILL_TOTAL   = PatternFill("solid", fgColor="D9E1F2")   # light blue

_FONT_WHITE_BOLD = Font(bold=True, color="FFFFFF")
_FONT_BOLD       = Font(bold=True)
_FONT_HEADER     = Font(bold=True, color="FFFFFF", size=11)

_THIN_BORDER = Border(
    left   = Side(style="thin"),
    right  = Side(style="thin"),
    top    = Side(style="thin"),
    bottom = Side(style="thin"),
)

_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center")


def _status_fill(status: str) -> PatternFill:
    return {"PASS": _FILL_PASS, "FAIL": _FILL_FAIL, "MISSING": _FILL_MISSING,
            "N/A": _FILL_NA}.get(
        status, _FILL_MISSING
    )


def _header_row(ws: Any, row: int, values: list[str], widths: list[int]) -> None:
    for col, (val, width) in enumerate(zip(values, widths), start=1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill   = _FILL_HEADER
        cell.font   = _FONT_HEADER
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER
        ws.column_dimensions[get_column_letter(col)].width = width


def _data_cell(ws: Any, row: int, col: int, value: str,
               fill: PatternFill | None = None,
               bold: bool = False,
               wrap: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = _THIN_BORDER
    if fill:
        cell.fill = fill
    font = Font(bold=bold)
    cell.font = font
    cell.alignment = _WRAP if wrap else _CENTER


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _sheet_summary(wb: Any, results: dict, tools: list[str], study: str,
                   study_dir: Path) -> None:
    ws = wb.active
    ws.title = "Summary"

    # Title row
    title = ws.cell(row=1, column=1, value=f"NGS-QC Execution Summary  |  {study}  |  {study_dir}")
    title.font = Font(bold=True, color="FFFFFF", size=12)
    title.fill = _FILL_TITLE
    title.alignment = _WRAP
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(tools) + 3)
    ws.row_dimensions[1].height = 24

    # Timestamp row
    ts = ws.cell(row=2, column=1,
                 value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ts.font = _FONT_BOLD
    ws.merge_cells(start_row=2, start_column=1,
                   end_row=2, end_column=len(tools) + 3)
    ws.row_dimensions[2].height = 18

    # Header row (row 3)
    headers = ["Sample"] + tools + ["Overall", "PASS count"]
    widths  = [28] + [14] * len(tools) + [12, 12]
    _header_row(ws, 3, headers, widths)

    samples = sorted(results.keys())
    pass_counts = {t: 0 for t in tools}
    fail_counts = {t: 0 for t in tools}

    for r_offset, sample in enumerate(samples):
        row = r_offset + 4
        _data_cell(ws, row, 1, sample, bold=True)
        tool_statuses = results[sample]
        all_pass = True
        for c_offset, tname in enumerate(tools):
            col     = c_offset + 2
            status  = tool_statuses[tname].status
            fill    = _status_fill(status)
            _data_cell(ws, row, col, status, fill=fill)
            if status == "PASS":
                pass_counts[tname] += 1
            elif status != "N/A":
                all_pass = False
                fail_counts[tname] += 1
        overall = "PASS" if all_pass else "FAIL"
        _data_cell(ws, row, len(tools) + 2, overall,
                   fill=_status_fill(overall), bold=True)
        applicable = [t for t in tools if tool_statuses[t].status != "N/A"]
        pass_n = sum(1 for t in applicable if tool_statuses[t].status == "PASS")
        _data_cell(ws, row, len(tools) + 3,
                   f"{pass_n}/{len(applicable)}")

    # Totals row
    tot_row = len(samples) + 4
    _data_cell(ws, tot_row, 1, "PASS count", fill=_FILL_TOTAL, bold=True)
    for c_offset, tname in enumerate(tools):
        _data_cell(ws, tot_row, c_offset + 2, str(pass_counts[tname]),
                   fill=_FILL_TOTAL, bold=True)

    fail_row = tot_row + 1
    _data_cell(ws, fail_row, 1, "FAIL / MISSING count", fill=_FILL_TOTAL, bold=True)
    for c_offset, tname in enumerate(tools):
        _data_cell(ws, fail_row, c_offset + 2, str(fail_counts[tname]),
                   fill=_FILL_TOTAL, bold=True)

    ws.freeze_panes = "B4"


def _sheet_failed(wb: Any, results: dict, tools: list[str]) -> None:
    ws = wb.create_sheet("Failed Jobs")
    headers = ["Sample", "Tool", "Status", "Log Path", "Fatal Error Lines"]
    widths  = [28, 15, 10, 60, 80]
    _header_row(ws, 1, headers, widths)

    row = 2
    for sample in sorted(results):
        for tname in tools:
            res = results[sample][tname]
            if res.status in ("PASS", "N/A"):
                continue
            _data_cell(ws, row, 1, sample, bold=True)
            _data_cell(ws, row, 2, tname)
            _data_cell(ws, row, 3, res.status, fill=_status_fill(res.status), bold=True)
            _data_cell(ws, row, 4, res.log_path, wrap=True)
            problems = [f"missing output: {p}" for p in res.missing_outputs] + res.fatal_lines[:10]
            snippet = "\n".join(problems) if problems else "(no fatal patterns detected)"
            _data_cell(ws, row, 5, snippet, wrap=True)
            ws.row_dimensions[row].height = max(30, min(15 * len(problems or [1]), 120))
            row += 1

    if row == 2:
        ws.cell(row=2, column=1, value="All tools PASSED for all samples.").font = _FONT_BOLD

    ws.freeze_panes = "A2"


def _sheet_tool(wb: Any, tool: str, results: dict) -> None:
    # Excel sheet names: max 31 chars, no []:*?/\\
    ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", tool)[:31])
    headers = ["Sample", "Status", "Output File(s)", "Log Path", f"Log tail (last {_LOG_TAIL_LINES} lines)"]
    widths  = [28, 10, 60, 60, 100]
    _header_row(ws, 1, headers, widths)

    for row, sample in enumerate(sorted(results), start=2):
        res = results[sample][tool]
        _data_cell(ws, row, 1, sample, bold=True)
        _data_cell(ws, row, 2, res.status, fill=_status_fill(res.status), bold=True)
        out_str = "\n".join(res.output_files) if res.output_files else "NOT FOUND"
        _data_cell(ws, row, 3, out_str, wrap=True)
        _data_cell(ws, row, 4, res.log_path, wrap=True)
        _data_cell(ws, row, 5, res.log_tail, wrap=True)
        ws.row_dimensions[row].height = 90

    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Report entry point
# ---------------------------------------------------------------------------

def write_report(
    study_dir: Path,
    study_name: str,
    output_path: Path,
) -> None:
    """Collect logs and write the multi-sheet XLSX report to *output_path*."""
    manifest = load_manifest(study_dir)
    tools    = manifest["tools"]
    results  = collect(manifest)
    n_samples = len(results)

    wb = openpyxl.Workbook()

    _sheet_summary(wb, results, tools, study_name, study_dir)
    _sheet_failed(wb, results, tools)
    for tool in tools:
        _sheet_tool(wb, tool, results)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    # --- console summary ---
    total_checks = sum(
        1 for sample_res in results.values()
        for res in sample_res.values() if res.status != "N/A"
    )
    n_pass = sum(
        1
        for sample_res in results.values()
        for res in sample_res.values()
        if res.status == "PASS"
    )
    n_fail = sum(
        1
        for sample_res in results.values()
        for res in sample_res.values()
        if res.status == "FAIL"
    )
    n_miss = total_checks - n_pass - n_fail  # N/A excluded from all counts

    print(f"[collect_qc_logs] Study   : {study_name}")
    print(f"[collect_qc_logs] Samples : {n_samples}")
    print(f"[collect_qc_logs] PASS    : {n_pass}/{total_checks}")
    print(f"[collect_qc_logs] FAIL    : {n_fail}/{total_checks}")
    print(f"[collect_qc_logs] MISSING : {n_miss}/{total_checks}")
    print(f"[collect_qc_logs] Report  : {output_path}")

    if n_fail > 0:
        print("\n[collect_qc_logs] Failed sample × tool combinations:")
        for sample in sorted(results):
            for tname, res in results[sample].items():
                if res.status == "FAIL":
                    print(f"  FAIL  {sample}  →  {tname}")
    if n_miss > 0:
        print("\n[collect_qc_logs] Missing (tool never ran) combinations:")
        for sample in sorted(results):
            for tname, res in results[sample].items():
                if res.status == "MISSING":
                    print(f"  MISS  {sample}  →  {tname}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="collect_qc_logs.py",
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "study_dir",
        metavar="STUDY_DIR",
        help="Per-study output directory created by run_ngs_qc.py.",
    )
    parser.add_argument(
        "study_name",
        metavar="STUDY_NAME",
        help="Study label used in the report title.",
    )
    parser.add_argument(
        "--output",
        metavar="XLSX",
        default=None,
        help=(
            "Path for the XLSX report. "
            "Defaults to STUDY_DIR/../STUDY_NAME_qc_execution_report.xlsx"
        ),
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args     = _parse_args(argv)
    study_dir = Path(args.study_dir).resolve()

    if not study_dir.is_dir():
        sys.exit(f"ERROR: STUDY_DIR does not exist or is not a directory: {study_dir}")

    default_out = study_dir.parent / f"{args.study_name}_qc_execution_report.xlsx"
    output_path = Path(args.output).resolve() if args.output else default_out

    write_report(study_dir, args.study_name, output_path)


if __name__ == "__main__":
    main()
